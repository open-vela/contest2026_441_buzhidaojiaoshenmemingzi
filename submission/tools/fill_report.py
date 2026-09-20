#!/usr/bin/env python3
"""Fill the official contest DOCX without changing its section structure."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "submission" / "2026首届openvela大赛作品提交模板.docx"
DRAFT = ROOT / "submission" / "技术报告正文草案.md"
OUTPUT = ROOT / "submission" / "技术报告-BK7258-R1-视觉辅助胸牌.docx"
DIAGRAM = ROOT / "submission" / "系统架构图.png"


def plain(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "").strip()


def markdown_sections(text: str) -> dict[str, str]:
    matches = list(re.finditer(r"(?m)^##\s+([^\n]+)\s*$", text))
    result: dict[str, str] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        key = match.group(1).strip()
        result[key] = text[match.end():end].strip()
    return result


def blocks(text: str):
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("### "):
            yield "heading", line[4:]
            index += 1
        elif line.startswith("!["):
            yield "image", line
            index += 1
        elif line.startswith("```"):
            index += 1
            code = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            yield "code", "\n".join(code)
            index += 1
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([plain(cell) for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            if len(rows) > 2 and all(re.fullmatch(r"[-: ]+", cell or " ") for cell in rows[1]):
                rows.pop(1)
            yield "table", rows
        elif re.match(r"^[-*]\s+|^\d+\.\s+", line):
            yield "list", re.sub(r"^[-*]\s+|^\d+\.\s+", "", line)
            index += 1
        else:
            paragraph = [line]
            index += 1
            while index < len(lines) and lines[index].strip() and not re.match(
                r"^(### |!\[|```|\||[-*]\s+|\d+\.\s+)", lines[index].strip()
            ):
                paragraph.append(lines[index].strip())
                index += 1
            yield "paragraph", " ".join(paragraph)


def set_paragraph_text(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(plain(text))
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def find_heading(paragraphs, prefix: str) -> int:
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip().startswith(prefix):
            return index
    raise RuntimeError(f"official template heading not found: {prefix}")


def insert_before(element, anchor) -> None:
    anchor._element.addprevious(element._element)


def add_body(document, anchor, kind: str, value) -> None:
    if kind == "table":
        rows = value
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index, cell_text in enumerate(row):
                cell = table.cell(row_index, column_index)
                cell.text = cell_text
                if row_index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "DCE7F2")
                    cell._tc.get_or_add_tcPr().append(shade)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        insert_before(table, anchor)
        return

    paragraph = document.add_paragraph()
    insert_before(paragraph, anchor)
    if kind == "heading":
        paragraph.style = "Heading 2"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_with_next = True
        set_paragraph_text(paragraph, value)
        paragraph.runs[0].bold = True
        paragraph.runs[0].font.size = Pt(11.5)
    elif kind == "list":
        paragraph.style = "List Bullet"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_paragraph_text(paragraph, value)
    elif kind == "code":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_text(paragraph, value)
        for run in paragraph.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
    elif kind == "image":
        match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", value)
        if not match:
            raise RuntimeError(f"invalid image reference: {value}")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(DIAGRAM), width=Inches(6.15))
        caption = document.add_paragraph()
        insert_before(caption, anchor)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_text(caption, match.group(1))
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.keep_together = True
        set_paragraph_text(paragraph, value)


def fill_section(document: Document, section: str, body: str,
                 next_section: str) -> None:
    paragraphs = document.paragraphs
    start = find_heading(paragraphs, section) + 1
    end = find_heading(paragraphs, next_section)
    slots = paragraphs[start:end]
    if not slots:
        raise RuntimeError(f"official template has no content slot for {section}")
    anchor = paragraphs[end]
    for paragraph in slots:
        paragraph._element.getparent().remove(paragraph._element)
    for kind, value in blocks(body):
        add_body(document, anchor, kind, value)


def main() -> None:
    document = Document(TEMPLATE)
    for name in ("Heading 1", "Heading 2", "List Bullet"):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    original_paragraphs = list(document.paragraphs)
    sections = markdown_sections(DRAFT.read_text(encoding="utf-8"))

    set_paragraph_text(document.paragraphs[0], "BK7258 R1 视觉辅助胸牌 OpenVela 技术报告")
    document.paragraphs[0].runs[0].font.size = Pt(20)
    document.paragraphs[0].runs[0].bold = True
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            paragraph.clear()  # Remove the template's unrelated picture watermark.

    info = document.tables[2]
    info.cell(1, 1).text = "BK7258 R1 视觉辅助胸牌 OpenVela"
    info.cell(2, 1).text = "不知道叫什么名字（队伍编号 441）"
    info.cell(3, 1).text = "邢晓亮（GitHub：xxluestc）：BSP 移植、驱动适配、构建烧录、实机验证、应用与报告整理"
    info.cell(4, 1).text = "新硬件平台适配"

    abstract = sections.get("2. 摘要") or sections.get("2、摘要")
    if not abstract:
        raise RuntimeError("draft abstract section missing")
    # The official template keeps the abstract instruction immediately after
    # paragraph 6 ("2、摘要").  Use the fixed template slot so the earlier
    # submission-checklist item numbered 2 cannot be mistaken for this title.
    set_paragraph_text(original_paragraphs[7], abstract)

    order = ["3.1", "3.2", "3.3", "3.4", "3.5", "3.6", "3.7"]
    for index, prefix in enumerate(order):
        body = next((value for key, value in sections.items()
                     if key.startswith(prefix)), None)
        if body is None:
            raise RuntimeError(f"draft section missing: {prefix}")
        next_prefix = order[index + 1] if index + 1 < len(order) else "四"
        fill_section(document, prefix, body, next_prefix)

    ai = document.tables[3]
    ai.cell(1, 1).text = "70%（团队估算，按本队新增和改动代码的开发参与度）"
    ai.cell(2, 1).text = "Codex、MiMo Code、本地终端、源码检索与文档处理"
    ai.cell(3, 1).text = "无"
    ai.cell(4, 1).text = "bk7258-openvela-porting；用于 BSP 分层、构建、实机验证和证据门禁，12 项测试通过"
    ai.cell(5, 1).text = "约 500,000 Token（团队估算，非平台账单统计）"

    ai_paragraphs = [paragraph for row in ai.rows for cell in row.cells
                     for paragraph in cell.paragraphs]
    for paragraph in ai_paragraphs[:-1]:
        paragraph.paragraph_format.keep_with_next = True

    for table in (info, ai):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "宋体"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                        run.font.size = Pt(10.5)

    # Keep the official report section, but remove the template's submission
    # instructions and scoring notes from the finished report.
    for paragraph in original_paragraphs[1:5] + original_paragraphs[38:]:
        if paragraph._element.getparent() is not None:
            paragraph._element.getparent().remove(paragraph._element)
    for table in (document.tables[0], document.tables[1], document.tables[4]):
        table._element.getparent().remove(table._element)

    for paragraph in document.paragraphs:
        if paragraph.text.strip() in ("3.1 绪论", "3.2 系统方案设计", "3.3 核心算法与技术原理",
                                      "3.4 系统实现", "3.5 系统测试与结果分析", "3.6 AI-Native 开发说明",
                                      "3.7 总结与展望"):
            paragraph.style = "Heading 1"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            if not run.font.name:
                run.font.name = "宋体"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")

    document.core_properties.title = "BK7258 R1 视觉辅助胸牌 OpenVela"
    document.core_properties.subject = "2026 首届 OpenVela AI 硬件开发者大赛技术报告"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
